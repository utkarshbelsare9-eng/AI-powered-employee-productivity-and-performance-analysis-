import os
from datetime import date
from flask import Blueprint, render_template, send_file, request, current_app, flash, redirect, url_for
from flask_login import login_required, current_user
import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.models import db
from app.models.employee import Employee
from app.models.user import User
from app.models.performance import Performance
from app.models.attendance import Attendance
from app.models.task import Task
from app.utils.helpers import role_required, current_month_str

reports_bp = Blueprint("reports", __name__)

BRAND_PURPLE = colors.HexColor("#6C5CE7")


@reports_bp.route("/reports")
@login_required
def reports_home():
    if current_user.is_employee:
        employees = [current_user.employee] if current_user.employee else []
    elif current_user.is_manager and current_user.employee:
        employees = Employee.query.filter_by(manager_id=current_user.employee.id).order_by(Employee.name).all()
    else:
        employees = Employee.query.outerjoin(User, Employee.user_id == User.id).filter(
            db.or_(User.role != "admin", User.id == None)
        ).order_by(Employee.name).all()
    return render_template("report.html", employees=employees)


@reports_bp.route("/reports/employee/<int:emp_id>/pdf")
@login_required
def employee_report_pdf(emp_id):
    emp = Employee.query.get_or_404(emp_id)
    if current_user.is_employee and (not current_user.employee or current_user.employee.id != emp.id):
        flash("You can only download your own report.", "danger")
        return redirect(url_for("reports.reports_home"))

    month = current_month_str()
    perf = Performance.query.filter_by(employee_id=emp.id, month=month).first()
    tasks = Task.query.filter_by(employee_id=emp.id).order_by(Task.deadline.desc()).limit(10).all()

    path = os.path.join(current_app.config["REPORTS_FOLDER"], f"employee_{emp.employee_code}_report.pdf")
    os.makedirs(current_app.config["REPORTS_FOLDER"], exist_ok=True)

    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleStyle", parent=styles["Title"], textColor=BRAND_PURPLE)
    heading_style = ParagraphStyle("H2", parent=styles["Heading2"], textColor=BRAND_PURPLE, spaceBefore=14)

    elements = [
        Paragraph("AI-Powered Employee Productivity &amp; Performance Analyzer", title_style),
        Paragraph(f"Employee Performance Report — {month}", styles["Heading3"]),
        Spacer(1, 12),
    ]

    info_data = [
        ["Employee Code", emp.employee_code, "Name", emp.name],
        ["Designation", emp.designation or "-", "Department", emp.department.name if emp.department else "-"],
        ["Status", emp.status, "Joining Date", str(emp.joining_date or "-")],
    ]
    info_table = Table(info_data, colWidths=[3.5 * cm, 4.5 * cm, 3.5 * cm, 4.5 * cm])
    info_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1EEFF")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F1EEFF")),
    ]))
    elements.append(info_table)

    elements.append(Paragraph("Performance Summary", heading_style))
    if perf:
        perf_data = [
            ["Metric", "Value"],
            ["Attendance %", f"{perf.attendance_percent}%"],
            ["Task Completion Rate", f"{perf.task_completion_rate}%"],
            ["Productivity Score", f"{perf.productivity_score}"],
            ["Performance Score", f"{perf.performance_score}"],
            ["Performance Grade", perf.performance_grade],
            ["Manager Rating", f"{perf.manager_rating} / 5"],
            ["Customer Rating", f"{perf.customer_rating} / 5"],
            ["Risk Score", f"{perf.risk_score}"],
        ]
        perf_table = Table(perf_data, colWidths=[8 * cm, 8 * cm])
        perf_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), BRAND_PURPLE),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F7FF")]),
        ]))
        elements.append(perf_table)
    else:
        elements.append(Paragraph("No performance record available for this month yet.", styles["Normal"]))

    elements.append(Paragraph("Recent Tasks", heading_style))
    if tasks:
        task_data = [["Title", "Status", "Priority", "Completed %", "Deadline"]]
        for t in tasks:
            task_data.append([t.title, t.status, t.priority, f"{t.completed_percent}%", str(t.deadline or "-")])
        task_table = Table(task_data, colWidths=[6 * cm, 2.8 * cm, 2.5 * cm, 2.7 * cm, 2 * cm])
        task_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), BRAND_PURPLE),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F7FF")]),
        ]))
        elements.append(task_table)
    else:
        elements.append(Paragraph("No tasks recorded.", styles["Normal"]))

    elements.append(Spacer(1, 20))
    elements.append(Paragraph(f"Generated on {date.today().isoformat()} by AI-Powered Employee Productivity & Performance Analyzer",
                               ParagraphStyle("Footer", parent=styles["Normal"], fontSize=7, textColor=colors.grey)))

    doc.build(elements)
    return send_file(path, as_attachment=True, download_name=f"{emp.employee_code}_report.pdf")


@reports_bp.route("/reports/employee/<int:emp_id>/excel")
@login_required
def employee_report_excel(emp_id):
    emp = Employee.query.get_or_404(emp_id)
    if current_user.is_employee and (not current_user.employee or current_user.employee.id != emp.id):
        flash("You can only download your own report.", "danger")
        return redirect(url_for("reports.reports_home"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Performance"

    header_fill = PatternFill(start_color="6C5CE7", end_color="6C5CE7", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append(["Employee Code", "Name", "Department", "Month", "Attendance %", "Task Completion %",
               "Productivity Score", "Performance Score", "Grade", "Manager Rating", "Customer Rating", "Risk Score"])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    perfs = Performance.query.filter_by(employee_id=emp.id).order_by(Performance.month).all()
    for p in perfs:
        ws.append([emp.employee_code, emp.name, emp.department.name if emp.department else "-", p.month,
                   p.attendance_percent, p.task_completion_rate, p.productivity_score,
                   p.performance_score, p.performance_grade, p.manager_rating, p.customer_rating, p.risk_score])

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 3

    path = os.path.join(current_app.config["REPORTS_FOLDER"], f"employee_{emp.employee_code}_report.xlsx")
    os.makedirs(current_app.config["REPORTS_FOLDER"], exist_ok=True)
    wb.save(path)
    return send_file(path, as_attachment=True, download_name=f"{emp.employee_code}_report.xlsx")


@reports_bp.route("/reports/company/excel")
@login_required
@role_required("admin", "manager")
def company_report_excel():
    month = current_month_str()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Performance {month}"

    header_fill = PatternFill(start_color="6C5CE7", end_color="6C5CE7", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(["Employee Code", "Name", "Department", "Designation", "Attendance %", "Task Completion %",
               "Productivity Score", "Performance Score", "Grade", "Status"])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    query = db.session.query(Employee, Performance).join(
        Performance, Performance.employee_id == Employee.id).outerjoin(
        User, Employee.user_id == User.id).filter(
        Performance.month == month,
        db.or_(User.role != "admin", User.id == None)
    )
    if current_user.is_manager and current_user.employee:
        query = query.filter(Employee.manager_id == current_user.employee.id)

    for emp, p in query.all():
        ws.append([emp.employee_code, emp.name, emp.department.name if emp.department else "-",
                   emp.designation, p.attendance_percent, p.task_completion_rate, p.productivity_score,
                   p.performance_score, p.performance_grade, emp.status])

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 3

    path = os.path.join(current_app.config["REPORTS_FOLDER"], f"company_report_{month}.xlsx")
    os.makedirs(current_app.config["REPORTS_FOLDER"], exist_ok=True)
    wb.save(path)
    return send_file(path, as_attachment=True, download_name=f"company_report_{month}.xlsx")
